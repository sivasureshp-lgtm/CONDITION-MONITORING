from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException, Form
from fastapi.responses import JSONResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timezone, timedelta
import chromadb
from chromadb.utils import embedding_functions
from sentence_transformers import SentenceTransformer
import PyPDF2
import io
from PIL import Image
import pytesseract
import openpyxl
import gspread
from google.oauth2.service_account import Credentials
from emergentintegrations.llm.chat import LlmChat, UserMessage
import json
import asyncio

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# ChromaDB setup
chroma_client = chromadb.PersistentClient(path="/app/backend/chroma_db")
embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

# Create or get collection
try:
    collection = chroma_client.get_collection(name="plant_knowledge")
except:
    collection = chroma_client.create_collection(
        name="plant_knowledge",
        metadata={"description": "Plant instrumentation knowledge base"}
    )

# LLM Setup
EMERGENT_KEY = os.environ.get('EMERGENT_LLM_KEY')

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# Models
class Document(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    filename: str
    doc_type: str  # Manual, SOP, Drawing, History
    content: str
    metadata: Dict[str, Any] = {}
    uploaded_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class ExpertQuery(BaseModel):
    query: str
    machine: Optional[str] = None
    line: Optional[str] = None
    severity: Optional[str] = None

class RagResponse(BaseModel):
    issue_summary: str
    key_observations: List[str]
    retrieved_knowledge: List[Dict[str, str]]
    root_cause_analysis: List[Dict[str, str]]
    recommended_actions: Dict[str, List[str]]
    drawing_reference: Dict[str, str]
    condition_monitoring: Dict[str, Any]
    confidence_level: str
    final_recommendation: str

class ConditionMonitoringData(BaseModel):
    plant: str  # A, G, K, E
    machine: str  # A1, A2, G1, K1, etc.
    motor: str  # Component name (e.g., TubeRotation, Sec1 Invert)
    current: float  # Measured current in Amps
    normal_current: float  # Normal threshold
    warning_current: float  # Warning threshold
    timestamp: datetime
    status: str  # OK, Warning, Alarm

class ConditionMonitoringCreate(BaseModel):
    plant: str
    machine: str
    motor: str
    current: float
    normal_current: float
    warning_current: float
    entry_source: str = "Office"  # Field or Office
    verified_by: Optional[str] = None
    notes: Optional[str] = None

# Helper Functions
def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF"""
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        logging.error(f"PDF extraction error: {e}")
        return ""

def extract_text_from_image(file_bytes: bytes) -> str:
    """Extract text from image using OCR"""
    try:
        image = Image.open(io.BytesIO(file_bytes))
        text = pytesseract.image_to_string(image)
        return text
    except Exception as e:
        logging.error(f"Image OCR error: {e}")
        return ""

def extract_text_from_excel(file_bytes: bytes) -> str:
    """Extract text from Excel"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell]) + "\n"
        return text
    except Exception as e:
        logging.error(f"Excel extraction error: {e}")
        return ""

async def generate_rag_response(query: str, context_docs: List[Dict], machine: str = None) -> RagResponse:
    """Generate structured RAG response using LLM"""
    
    # Prepare context
    context = "\n\n".join([
        f"Source: {doc['source']}\nDocument: {doc['document']}\nContent: {doc['content']}"
        for doc in context_docs
    ])
    
    system_prompt = """You are a Senior Instrumentation & Process Control Expert with 20+ years of experience.
Your role is to provide expert troubleshooting guidance for industrial plants.
Always structure your response as a valid JSON object with the following fields:
- issue_summary (string)
- key_observations (array of strings)
- retrieved_knowledge (array of objects with source, document, section, key_extract)
- root_cause_analysis (array of objects with cause and justification)
- recommended_actions (object with immediate, detailed_troubleshooting, preventive arrays)
- drawing_reference (object with drawing_type and what_to_verify)
- condition_monitoring (object with parameters_to_verify, trend_to_observe, sheet_update_required)
- confidence_level (string: High/Medium/Low)
- final_recommendation (string)

Be technical, precise, and practical. Prioritize plant safety."""
    
    user_prompt = f"""Query: {query}
Machine/Line: {machine or 'Not specified'}

Relevant Knowledge Base:
{context}

Provide a comprehensive troubleshooting analysis in JSON format."""
    
    try:
        chat = LlmChat(
            api_key=EMERGENT_KEY,
            session_id=f"expert_{uuid.uuid4()}",
            system_message=system_prompt
        ).with_model("openai", "gpt-4o")
        
        message = UserMessage(text=user_prompt)
        response = await chat.send_message(message)
        
        # Parse JSON response
        response_json = json.loads(response)
        return RagResponse(**response_json)
    
    except Exception as e:
        logging.error(f"LLM generation error: {e}")
        # Return fallback response
        return RagResponse(
            issue_summary=f"Issue: {query}",
            key_observations=["Unable to generate detailed analysis"],
            retrieved_knowledge=context_docs[:3],
            root_cause_analysis=[{"cause": "Analysis in progress", "justification": "Please retry"}],
            recommended_actions={
                "immediate": ["Contact maintenance team"],
                "detailed_troubleshooting": ["Review relevant manuals"],
                "preventive": ["Schedule preventive maintenance"]
            },
            drawing_reference={"drawing_type": "P&ID", "what_to_verify": "Check instrument loops"},
            condition_monitoring={"parameters_to_verify": ["Temperature", "Pressure"], "trend_to_observe": "Check recent trends", "sheet_update_required": True},
            confidence_level="Low",
            final_recommendation="Requires further investigation with complete data"
        )

# API Routes
@api_router.get("/")
async def root():
    return {"message": "Process Control Expert System API"}

@api_router.post("/documents/upload")
async def upload_document(
    file: UploadFile = File(...),
    doc_type: str = Form(...),
    machine: str = Form(None),
    section: str = Form(None)
):
    """Upload and process document"""
    try:
        file_bytes = await file.read()
        
        # Extract text based on file type
        if file.filename.endswith('.pdf'):
            text = extract_text_from_pdf(file_bytes)
        elif file.filename.endswith(('.png', '.jpg', '.jpeg')):
            text = extract_text_from_image(file_bytes)
        elif file.filename.endswith(('.xlsx', '.xls')):
            text = extract_text_from_excel(file_bytes)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")
        
        if not text.strip():
            raise HTTPException(status_code=400, detail="No text extracted from document")
        
        # Create document
        doc_id = str(uuid.uuid4())
        doc = Document(
            id=doc_id,
            filename=file.filename,
            doc_type=doc_type,
            content=text[:5000],  # Store preview
            metadata={
                "machine": machine,
                "section": section,
                "file_size": len(file_bytes)
            }
        )
        
        # Store in MongoDB
        doc_dict = doc.model_dump()
        doc_dict['uploaded_at'] = doc_dict['uploaded_at'].isoformat()
        await db.documents.insert_one(doc_dict)
        
        # Add to vector store
        # Split text into chunks for better retrieval
        chunk_size = 1000
        chunks = [text[i:i+chunk_size] for i in range(0, len(text), chunk_size)]
        
        for idx, chunk in enumerate(chunks[:10]):  # Limit to 10 chunks per doc
            chunk_id = f"{doc_id}_{idx}"
            collection.add(
                ids=[chunk_id],
                documents=[chunk],
                metadatas=[{
                    "doc_id": doc_id,
                    "filename": file.filename,
                    "doc_type": doc_type,
                    "machine": machine or "general",
                    "section": section or "general",
                    "chunk_index": idx
                }]
            )
        
        return {"message": "Document uploaded successfully", "doc_id": doc_id, "chunks": len(chunks[:10])}
    
    except Exception as e:
        logging.error(f"Upload error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@api_router.get("/documents")
async def get_documents():
    """Get all documents"""
    docs = await db.documents.find({}, {"_id": 0}).to_list(1000)
    for doc in docs:
        if isinstance(doc.get('uploaded_at'), str):
            doc['uploaded_at'] = datetime.fromisoformat(doc['uploaded_at'])
    return docs

@api_router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str):
    """Delete document"""
    # Delete from MongoDB
    await db.documents.delete_one({"id": doc_id})
    
    # Delete from vector store
    try:
        # Get all chunk IDs for this document
        results = collection.get(where={"doc_id": doc_id})
        if results['ids']:
            collection.delete(ids=results['ids'])
    except Exception as e:
        logging.error(f"Vector delete error: {e}")
    
    return {"message": "Document deleted successfully"}

@api_router.post("/query", response_model=RagResponse)
async def expert_query(query: ExpertQuery):
    """Process expert query with RAG"""
    try:
        # Retrieve relevant documents from vector store
        results = collection.query(
            query_texts=[query.query],
            n_results=5,
            where={"machine": query.machine} if query.machine else None
        )
        
        # Format retrieved knowledge
        context_docs = []
        if results['documents'] and results['documents'][0]:
            for idx, doc in enumerate(results['documents'][0]):
                metadata = results['metadatas'][0][idx]
                context_docs.append({
                    "source": metadata.get('doc_type', 'Unknown'),
                    "document": metadata.get('filename', 'Unknown'),
                    "section": metadata.get('section', 'N/A'),
                    "content": doc[:500]
                })
        
        # Generate RAG response
        response = await generate_rag_response(query.query, context_docs, query.machine)
        
        # Store query history
        query_record = {
            "id": str(uuid.uuid4()),
            "query": query.query,
            "machine": query.machine,
            "line": query.line,
            "severity": query.severity,
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "confidence": response.confidence_level
        }
        await db.query_history.insert_one(query_record)
        
        return response
    
    except Exception as e:
        logging.error(f"Query error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@api_router.get("/query/history")
async def get_query_history():
    """Get query history"""
    history = await db.query_history.find({}, {"_id": 0}).sort("timestamp", -1).to_list(50)
    return history

@api_router.post("/condition-monitoring")
async def add_condition_data(data: ConditionMonitoringCreate):
    """Add condition monitoring data"""
    # Calculate status based on thresholds
    status = "OK"
    if data.current >= data.warning_current:
        status = "Alarm"
    elif data.current >= data.normal_current:
        status = "Warning"
    
    reading_timestamp = datetime.now(timezone.utc)
    entry_timestamp = datetime.now(timezone.utc)
    
    # Check for suspicious bulk entry (multiple readings within 1 minute)
    recent_entries = await db.condition_monitoring.count_documents({
        "plant": data.plant,
        "entry_timestamp": {
            "$gte": (entry_timestamp - timedelta(minutes=1)).isoformat()
        }
    })
    
    bulk_entry_flag = recent_entries > 5  # Flag if more than 5 entries in 1 minute
    
    doc = {
        "plant": data.plant,
        "machine": data.machine,
        "motor": data.motor,
        "current": data.current,
        "normal_current": data.normal_current,
        "warning_current": data.warning_current,
        "status": status,
        "timestamp": reading_timestamp.isoformat(),
        "entry_timestamp": entry_timestamp.isoformat(),
        "entry_source": data.entry_source,
        "verified_by": data.verified_by,
        "notes": data.notes,
        "bulk_entry_flag": bulk_entry_flag,
        "verified": data.entry_source == "Field"  # Field entries auto-verified
    }
    await db.condition_monitoring.insert_one(doc)
    return {
        "message": "Data added successfully", 
        "status": status,
        "bulk_entry_flag": bulk_entry_flag
    }

@api_router.get("/condition-monitoring/plant/{plant}")
async def get_plant_data(plant: str, limit: int = 1000):
    """Get condition monitoring data for a plant"""
    data = await db.condition_monitoring.find(
        {"plant": plant},
        {"_id": 0}
    ).sort("timestamp", -1).to_list(limit)
    return data

@api_router.get("/condition-monitoring/machine/{plant}/{machine}")
async def get_machine_data(plant: str, machine: str, limit: int = 100):
    """Get condition monitoring data for a specific machine"""
    data = await db.condition_monitoring.find(
        {"plant": plant, "machine": machine},
        {"_id": 0}
    ).sort("timestamp", -1).to_list(limit)
    return data

@api_router.get("/active-alarms")
async def get_active_alarms():
    """Get all active alarms"""
    pipeline = [
        {"$sort": {"timestamp": -1}},
        {"$group": {
            "_id": {"plant": "$plant", "machine": "$machine", "motor": "$motor"},
            "latest": {"$first": "$$ROOT"}
        }},
        {"$replaceRoot": {"newRoot": "$latest"}},
        {"$match": {"status": "Alarm"}},
        {"$project": {"_id": 0}}
    ]
    alarms = await db.condition_monitoring.aggregate(pipeline).to_list(100)
    return alarms

@api_router.get("/machine-health/{plant}")
async def get_machine_health(plant: str):
    """Get health status for all machines in a plant"""
    pipeline = [
        {"$match": {"plant": plant}},
        {"$sort": {"timestamp": -1}},
        {"$group": {
            "_id": {"plant": "$plant", "machine": "$machine", "motor": "$motor"},
            "latest": {"$first": "$$ROOT"}
        }},
        {"$replaceRoot": {"newRoot": "$latest"}},
        {"$group": {
            "_id": "$machine",
            "ok_count": {"$sum": {"$cond": [{"$eq": ["$status", "OK"]}, 1, 0]}},
            "warning_count": {"$sum": {"$cond": [{"$eq": ["$status", "Warning"]}, 1, 0]}},
            "alarm_count": {"$sum": {"$cond": [{"$eq": ["$status", "Alarm"]}, 1, 0]}},
            "total": {"$sum": 1}
        }},
        {"$project": {
            "_id": 0,
            "machine": "$_id",
            "ok": "$ok_count",
            "warning": "$warning_count",
            "alarm": "$alarm_count",
            "total": "$total",
            "health_percent": {
                "$round": [
                    {"$multiply": [
                        {"$divide": ["$ok_count", "$total"]},
                        100
                    ]},
                    0
                ]
            }
        }},
        {"$sort": {"machine": 1}}
    ]
    health_data = await db.condition_monitoring.aggregate(pipeline).to_list(100)
    return health_data

@api_router.get("/plant-health")
async def get_plant_health():
    """Get overall health for all plants"""
    pipeline = [
        {"$sort": {"timestamp": -1}},
        {"$group": {
            "_id": {"plant": "$plant", "machine": "$machine", "motor": "$motor"},
            "latest": {"$first": "$$ROOT"}
        }},
        {"$replaceRoot": {"newRoot": "$latest"}},
        {"$group": {
            "_id": "$plant",
            "ok_count": {"$sum": {"$cond": [{"$eq": ["$status", "OK"]}, 1, 0]}},
            "warning_count": {"$sum": {"$cond": [{"$eq": ["$status", "Warning"]}, 1, 0]}},
            "alarm_count": {"$sum": {"$cond": [{"$eq": ["$status", "Alarm"]}, 1, 0]}},
            "total": {"$sum": 1}
        }},
        {"$project": {
            "_id": 0,
            "plant": "$_id",
            "ok": "$ok_count",
            "warning": "$warning_count",
            "alarm": "$alarm_count",
            "total": "$total",
            "health_percent": {
                "$round": [
                    {"$multiply": [
                        {"$divide": ["$ok_count", "$total"]},
                        100
                    ]},
                    0
                ]
            }
        }},
        {"$sort": {"plant": 1}}
    ]
    health_data = await db.condition_monitoring.aggregate(pipeline).to_list(100)
    return health_data

@api_router.get("/stats")
async def get_stats():
    """Get system stats"""
    doc_count = await db.documents.count_documents({})
    query_count = await db.query_history.count_documents({})
    
    # Get document type breakdown
    pipeline = [
        {"$group": {"_id": "$doc_type", "count": {"$sum": 1}}}
    ]
    doc_types = await db.documents.aggregate(pipeline).to_list(100)
    
    return {
        "total_documents": doc_count,
        "total_queries": query_count,
        "document_types": {item['_id']: item['count'] for item in doc_types},
        "vector_store_size": collection.count()
    }

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()